"""
google_drive_integration.py — เชื่อมต่อ Google Drive จริง (อ่านอย่างเดียว) สำหรับฐานความรู้สินค้า
======================================================================
ใช้ OAuth 2.0 แบบ "Desktop app" — ผู้ใช้ต้องสร้าง Google Cloud Project + เปิด Drive API +
สร้าง OAuth Client ID (Desktop app) ด้วยตัวเอง แล้วโหลดไฟล์ credentials.json มาวางไว้ข้าง ๆ
ไฟล์นี้ (ดูคู่มือขั้นตอนที่ Claude อธิบายไว้ในแชท) — โค้ดนี้จะไม่สร้างบัญชี Google ใด ๆ
และจะไม่กดยืนยัน/อนุญาต (consent) แทนผู้ใช้เด็ดขาด ผู้ใช้ต้อง login + กด "อนุญาต" ด้วยตัวเองเสมอ

การทำงาน:
- กดปุ่ม "เชื่อมต่อ Google Drive" ครั้งแรก → เปิดเบราว์เซอร์ในเครื่องที่รันแอปนี้ ให้ผู้ใช้ login +
  อนุญาตสิทธิ์ "อ่านไฟล์อย่างเดียว" (drive.readonly) เอง
- หลังเชื่อมต่อสำเร็จ จะบันทึก token ไว้ที่ drive_token.json เพื่อใช้ซ้ำ (รีเฟรชอัตโนมัติ ไม่ต้อง
  login ใหม่ทุกครั้ง) จนกว่าจะกด "ยกเลิกการเชื่อมต่อ"
- ฟังก์ชัน sync_folder(folder_id) จะโหลดไฟล์ทั้งหมดในโฟลเดอร์ (PDF / Word / Excel / ข้อความ /
  Google Docs / Google Sheets) แตกเป็นข้อความ ให้หน้าเว็บนำไปบันทึกต่อใน product_knowledge.json
"""

import os
import re

_THIS_DIR = os.path.dirname(os.path.abspath(__file__))
CREDENTIALS_FILE = os.path.join(_THIS_DIR, "credentials.json")
TOKEN_FILE = os.path.join(_THIS_DIR, "drive_token.json")

# ขอสิทธิ์ "อ่านอย่างเดียว" เท่านั้น — ไม่ขอสิทธิ์แก้ไข/ลบ/อัปโหลดไฟล์ใน Drive ของผู้ใช้
SCOPES = ["https://www.googleapis.com/auth/drive.readonly"]

GOOGLE_DOC_MIME = "application/vnd.google-apps.document"
GOOGLE_SHEET_MIME = "application/vnd.google-apps.spreadsheet"
GOOGLE_FOLDER_MIME = "application/vnd.google-apps.folder"


def is_configured() -> bool:
    """เช็คว่าผู้ใช้วางไฟล์ credentials.json (ที่ดาวน์โหลดจาก Google Cloud Console เอง) ไว้แล้วหรือยัง"""
    return os.path.exists(CREDENTIALS_FILE)


def is_connected() -> bool:
    """เช็คว่าเคยทำ OAuth login จนได้ token บันทึกไว้แล้วหรือยัง"""
    return os.path.exists(TOKEN_FILE)


def libraries_installed() -> bool:
    try:
        import google.auth  # noqa: F401
        import googleapiclient  # noqa: F401
        import google_auth_oauthlib  # noqa: F401
        return True
    except ImportError:
        return False


def _get_credentials():
    """คืน Credentials ที่ใช้งานได้ (รีเฟรชอัตโนมัติถ้าหมดอายุ) หรือ None ถ้ายังไม่เชื่อมต่อ/หมดอายุแบบกู้ไม่ได้"""
    from google.oauth2.credentials import Credentials
    from google.auth.transport.requests import Request

    if not os.path.exists(TOKEN_FILE):
        return None
    creds = Credentials.from_authorized_user_file(TOKEN_FILE, SCOPES)
    if creds and creds.expired and creds.refresh_token:
        try:
            creds.refresh(Request())
            with open(TOKEN_FILE, "w", encoding="utf-8") as f:
                f.write(creds.to_json())
        except Exception:
            return None
    return creds


def connect() -> tuple:
    """เริ่ม OAuth flow จริง — เปิดเบราว์เซอร์ของเครื่องนี้ให้ผู้ใช้ login + กด "อนุญาต" ด้วยตัวเอง
    (สิทธิ์อ่านอย่างเดียว) ต้องรันบนเครื่องเดียวกับที่เปิดแอป Streamlit อยู่ (ไม่ใช่เซิร์ฟเวอร์ระยะไกล)
    คืน (สำเร็จหรือไม่: bool, ข้อความ: str)"""
    if not is_configured():
        return False, "ไม่พบไฟล์ credentials.json — กรุณาทำตามขั้นตอนสร้าง Google Cloud Project ก่อน (ดูคำแนะนำด้านบน)"
    if not libraries_installed():
        return False, "ยังไม่ได้ติดตั้งไลบรารี Google — รัน: pip install google-api-python-client google-auth-httplib2 google-auth-oauthlib"
    try:
        from google_auth_oauthlib.flow import InstalledAppFlow
        flow = InstalledAppFlow.from_client_secrets_file(CREDENTIALS_FILE, SCOPES)
        creds = flow.run_local_server(port=0)
        with open(TOKEN_FILE, "w", encoding="utf-8") as f:
            f.write(creds.to_json())
        return True, "✅ เชื่อมต่อ Google Drive สำเร็จ (สิทธิ์อ่านอย่างเดียว)"
    except Exception as e:
        return False, f"⚠️ เชื่อมต่อไม่สำเร็จ: {str(e)[:300]}"


def disconnect() -> None:
    """ลบ token ที่บันทึกไว้ — ครั้งถัดไปต้อง login ใหม่"""
    if os.path.exists(TOKEN_FILE):
        os.remove(TOKEN_FILE)


def extract_folder_id(url_or_id: str) -> str:
    """รับได้ทั้งลิงก์โฟลเดอร์ Drive แบบเต็ม หรือ folder id ตรง ๆ — คืน folder id เปล่า ๆ"""
    if not url_or_id:
        return ""
    m = re.search(r"/folders/([a-zA-Z0-9_-]+)", url_or_id)
    if m:
        return m.group(1)
    return url_or_id.strip()


def list_files_in_folder(folder_id: str) -> list:
    """คืน [{id, name, mimeType}, ...] ของไฟล์ทั้งหมดในโฟลเดอร์ (ไม่ลงโฟลเดอร์ย่อย)"""
    from googleapiclient.discovery import build
    creds = _get_credentials()
    if not creds:
        return []
    service = build("drive", "v3", credentials=creds)
    files, page_token = [], None
    while True:
        resp = service.files().list(
            q=f"'{folder_id}' in parents and trashed = false",
            fields="nextPageToken, files(id, name, mimeType)",
            pageToken=page_token,
        ).execute()
        files.extend(resp.get("files", []))
        page_token = resp.get("nextPageToken")
        if not page_token:
            break
    return files


def _extract_text_from_drive_file(service, f: dict) -> str:
    """ดาวน์โหลด/แปลงไฟล์ Drive 1 ไฟล์ ให้กลายเป็นข้อความล้วน"""
    import pdfplumber
    from docx import Document
    from io import BytesIO
    from googleapiclient.http import MediaIoBaseDownload

    mime = f.get("mimeType", "")
    name = f.get("name", "")

    try:
        if mime == GOOGLE_DOC_MIME:
            data = service.files().export(fileId=f["id"], mimeType="text/plain").execute()
            return data.decode("utf-8", errors="replace") if isinstance(data, bytes) else str(data)
        if mime == GOOGLE_SHEET_MIME:
            data = service.files().export(fileId=f["id"], mimeType="text/csv").execute()
            return data.decode("utf-8", errors="replace") if isinstance(data, bytes) else str(data)

        # ไฟล์ไบนารีปกติ (pdf/docx/xlsx/txt ฯลฯ) — ดาวน์โหลดเนื้อไฟล์มาก่อน
        request = service.files().get_media(fileId=f["id"])
        buf = BytesIO()
        downloader = MediaIoBaseDownload(buf, request)
        done = False
        while not done:
            _, done = downloader.next_chunk()
        buf.seek(0)

        ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
        if ext == "pdf" or mime == "application/pdf":
            with pdfplumber.open(buf) as pdf:
                return "\n".join(p.extract_text() or "" for p in pdf.pages)
        elif ext == "docx":
            doc = Document(buf)
            return "\n".join(p.text for p in doc.paragraphs)
        elif ext == "xlsx":
            try:
                import openpyxl
                wb = openpyxl.load_workbook(buf, data_only=True)
                rows = []
                for ws in wb.worksheets:
                    for row in ws.iter_rows(values_only=True):
                        rows.append(" | ".join(str(c) for c in row if c is not None))
                return "\n".join(rows)
            except ImportError:
                return "[ต้องติดตั้ง openpyxl เพื่ออ่านไฟล์ Excel นี้ — pip install openpyxl]"
        elif ext == "txt" or mime.startswith("text/"):
            return buf.read().decode("utf-8", errors="replace")
        else:
            return f"[ไม่รองรับไฟล์ประเภทนี้: {mime}]"
    except Exception as e:
        return f"[อ่านไฟล์ {name} ไม่สำเร็จ: {str(e)[:150]}]"


def sync_folder(folder_id: str) -> list:
    """ดึงไฟล์ทั้งหมดในโฟลเดอร์ Drive นี้ + แตกข้อความ — คืน [{filename, text}, ...]
    (ไม่ลงโฟลเดอร์ย่อย — sync เฉพาะไฟล์ในชั้นบนสุดของโฟลเดอร์ที่ระบุก่อน)"""
    from googleapiclient.discovery import build
    creds = _get_credentials()
    if not creds:
        return []
    service = build("drive", "v3", credentials=creds)
    out = []
    for f in list_files_in_folder(folder_id):
        if f.get("mimeType") == GOOGLE_FOLDER_MIME:
            continue
        out.append({"filename": f.get("name", "untitled"), "text": _extract_text_from_drive_file(service, f)})
    return out
